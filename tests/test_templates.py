"""
Tests for the Excel templates shipped with the package.
"""

import pytest

pytest.importorskip("openpyxl")

from actuneo.templates import list_templates, template_path, save_template  # noqa: E402
from actuneo.loss_reserving import Triangle, ChainLadder  # noqa: E402
from actuneo.ifrs17 import load_paa_case, example_case_path  # noqa: E402


def test_templates_are_shipped():
    assert list_templates() == ["paa_case", "triangle"]
    for name in list_templates():
        assert template_path(name).exists()
    with pytest.raises(ValueError, match="Unknown template"):
        template_path("pensions")


def test_triangle_template_loads_and_projects(tmp_path):
    path = save_template("triangle", str(tmp_path))
    assert path.name == "triangle_template.xlsx"
    triangle = Triangle.from_excel(str(path))
    assert triangle.shape == (5, 5)
    assert triangle.origin == [2021, 2022, 2023, 2024, 2025]
    assert ChainLadder(triangle).total_ibnr > 0


def test_paa_template_matches_the_example_case(tmp_path):
    path = save_template("paa_case", str(tmp_path / "my_case.xlsx"))
    assert path.name == "my_case.xlsx"
    from_template = load_paa_case(str(path))
    from_csv = load_paa_case(example_case_path())
    assert list(from_template.groups) == list(from_csv.groups)
    assert from_template.statements.balance_check() == pytest.approx(0, abs=1e-6)
    profit = "Profit for the period"
    assert from_template.statements.profit_or_loss().loc[profit].sum() == pytest.approx(
        from_csv.statements.profit_or_loss().loc[profit].sum()
    )


def test_templates_have_a_guide_sheet():
    import openpyxl
    for name in list_templates():
        book = openpyxl.load_workbook(template_path(name))
        assert "Guide" in book.sheetnames
        assert book.sheetnames[0] != "Guide"


def test_existing_file_is_not_overwritten(tmp_path):
    save_template("triangle", str(tmp_path))
    with pytest.raises(FileExistsError):
        save_template("triangle", str(tmp_path))
    assert save_template("triangle", str(tmp_path), overwrite=True).exists()
